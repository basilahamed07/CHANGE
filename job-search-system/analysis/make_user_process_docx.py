"""make_user_process_docx.py — render docs/<name>.md into a styled Word document.

Output: docs/<name>.docx  (open in Word / Google Docs / LibreOffice)
Run:    cd job-search-system/root && .venv/bin/python ../analysis/make_user_process_docx.py [name]

    name defaults to USER_PROCESS; pass any docs/<name>.md, e.g.:
    .venv/bin/python ../analysis/make_user_process_docx.py DAILY_RUN_GUIDE

Parse rules cover the docs' shared structure:
headings (#..####), tables, bullet lists, numbered lists, ``` code fences,
--- rules, **bold** / `code` inline runs.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Cm

DOCS = Path(__file__).resolve().parent.parent / "docs"
NAME = sys.argv[1] if len(sys.argv) > 1 else "USER_PROCESS"
MD = DOCS / f"{NAME}.md"
OUT = DOCS / f"{NAME}.docx"

INK = RGBColor(0x1C, 0x24, 0x30)
ACCENT = RGBColor(0x0B, 0x62, 0xD6)
GREY = RGBColor(0x64, 0x74, 0x8B)


def _shade(cell, hex_fill: str) -> None:
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _style_doc(doc: Document) -> None:
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.font.color.rgb = INK

    for name, size, color, before, after in [
        ("Heading 1", 20, ACCENT, 6, 10),
        ("Heading 2", 15, ACCENT, 14, 6),
        ("Heading 3", 12.5, INK, 10, 4),
        ("Heading 4", 11, GREY, 8, 3),
    ]:
        h = doc.styles[name]
        h.font.name = "Calibri"
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = color
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(after)


def _add_runs(par, text: str) -> None:
    """Inline markdown: **bold**, `code`, [text](url) → styled runs."""
    tokens = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(tok[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x0F, 0x4C, 0xA8)
        elif tok.startswith("[") and "](" in tok:
            label = tok[1:tok.index("](")]
            par.add_run(label).font.underline = True  # printed docs: no links
        else:
            par.add_run(tok)


def _add_table(doc: Document, rows: list[list[str]]) -> None:
    n_cols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=n_cols)
    t.style = "Light Grid Accent 1"
    for i, row in enumerate(rows):
        for j in range(n_cols):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            _add_runs(p, row[j] if j < len(row) else "")
            for r in p.runs:
                r.font.size = Pt(9.5)
                if i == 0:
                    r.font.bold = True
        if i == 0:
            for c in t.rows[0].cells:
                _shade(c, "E8F0FE")
    doc.add_paragraph()


def convert() -> Path:
    text = MD.read_text(encoding="utf-8")
    doc = Document()
    _style_doc(doc)

    lines = text.splitlines()
    i, table_buf, in_code = 0, [], False
    code_buf: list[str] = []

    def flush_table():
        nonlocal table_buf
        if table_buf:
            _add_table(doc, table_buf)
            table_buf = []

    def flush_code():
        nonlocal code_buf
        if code_buf:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4)
            pPr = p._p.get_or_add_pPr()
            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "F4F6F8")
            pPr.append(shd)
            for ln in code_buf:
                r = p.add_run(ln)
                r.font.name = "Consolas"
                r.font.size = Pt(9)
            doc.add_paragraph()
            code_buf = []

    while i < len(lines):
        line = lines[i]

        if line.startswith("```"):
            flush_table()
            if in_code:
                flush_code()
            in_code = not in_code
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-+:?", c) for c in cells):
                i += 1
                continue
            table_buf.append(cells)
            i += 1
            continue
        flush_table()

        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            lvl = min(len(m.group(1)) + 0, 4)
            doc.add_heading(m.group(2), level=lvl)
            i += 1
            continue

        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            _add_runs(p, line[2:])
            for r in p.runs:
                r.font.color.rgb = GREY
                r.italic = True
            i += 1
            continue

        if line.strip() in ("---", "***", "___"):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement
            pBdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), "CBD5E1")
            pBdr.append(bottom)
            pPr.append(pBdr)
            i += 1
            continue

        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, line[2:])
            i += 1
            continue

        mnum = re.match(r"^(\d+)\.\s+(.*)$", line)
        if mnum:
            p = doc.add_paragraph(style="List Number")
            _add_runs(p, mnum.group(2))
            i += 1
            continue

        if line.strip():
            p = doc.add_paragraph()
            _add_runs(p, line)
        i += 1

    flush_table()
    flush_code()
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    out = convert()
    print(f"wrote {out}")
    sys.exit(0)
