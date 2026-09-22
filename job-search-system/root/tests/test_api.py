import json
from unittest.mock import AsyncMock, MagicMock

import pytest
from httpx import AsyncClient, ASGITransport

from app.database import Database


@pytest.fixture
async def app(tmp_path):
    from app.main import create_app
    application = create_app(db_path=str(tmp_path / "test.db"), testing=True)
    db = Database(str(tmp_path / "test.db"))
    await db.init()
    application.state.db = db
    application.state.embedding_client = None
    yield application
    await db.close()


@pytest.fixture
async def client(app):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        yield ac


@pytest.mark.asyncio
async def test_health(client):
    resp = await client.get("/api/health")
    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "healthy"
    assert data["db"] == "ok"
    assert data["scheduler"] in ("running", "stopped", "not_configured")
    assert "ai_provider" in data
    assert "ai_configured" in data
    assert "last_scrape" in data
    assert "uptime_seconds" in data


@pytest.mark.asyncio
async def test_list_jobs_empty(client):
    resp = await client.get("/api/jobs")
    assert resp.status_code == 200
    assert resp.json()["jobs"] == []


@pytest.mark.asyncio
async def test_get_stats(client):
    resp = await client.get("/api/stats")
    assert resp.status_code == 200
    data = resp.json()
    assert "total_jobs" in data
    assert "total_scored" in data
    assert "total_applied" in data


@pytest.mark.asyncio
async def test_get_job_not_found(client):
    resp = await client.get("/api/jobs/999")
    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_trigger_scrape(client):
    resp = await client.post("/api/scrape")
    assert resp.status_code == 202
    body = resp.json()
    assert body["status"] == "started"
    assert isinstance(body["task_id"], str) and body["task_id"]


@pytest.mark.asyncio
async def test_dismiss_job_not_found(client):
    resp = await client.post("/api/jobs/999/dismiss")
    assert resp.status_code in [200, 404]


@pytest.mark.asyncio
async def test_prepare_not_found(client):
    resp = await client.post("/api/jobs/999/prepare")
    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_prepare_no_tailor(client, app):
    db = app.state.db
    job_id = await db.insert_job(
        title="Engineer", company="Acme", location="Remote",
        salary_min=150000, salary_max=200000,
        description="Build things", url="https://example.com/job1",
        posted_date="2026-01-01", application_method="url",
        contact_email=None,
    )
    app.state.tailor = None
    resp = await client.post(f"/api/jobs/{job_id}/prepare")
    assert resp.status_code == 503


@pytest.mark.asyncio
async def test_prepare_with_mock_tailor(client, app, tmp_path):
    db = app.state.db
    job_id = await db.insert_job(
        title="Engineer", company="Acme", location="Remote",
        salary_min=150000, salary_max=200000,
        description="Build things", url="https://example.com/job2",
        posted_date="2026-01-01", application_method="url",
        contact_email=None,
    )
    await db.insert_score(
        job_id, 85, ["Good skills match"], ["No concerns"], ["Python", "AWS"],
    )

    mock_tailor = MagicMock()
    mock_tailor.prepare = AsyncMock(return_value={
        "tailored_resume": "Tailored resume text",
        "cover_letter": "Dear hiring manager...",
    })
    app.state.tailor = mock_tailor

    # M2: the prepare endpoint is fail-closed behind the EvidenceChecker gate.
    # Provide a minimal verified corpus covering the mocked generation output so
    # this endpoint-mechanics test passes the gate (gate semantics have their
    # own dedicated tests in test_evidence.py).
    from app.evidence_store import EvidenceStore
    from app.evidence_checker import EvidenceChecker

    profile_dir = tmp_path / "profile"
    profile_dir.mkdir()
    (profile_dir / "experience.yaml").write_text(
        "claims:\n"
        "  - id: r1\n"
        "    value: \"Tailored resume text\"\n"
        "    status: VERIFIED\n"
        "    source: test\n"
        "  - id: c1\n"
        "    value: \"Dear hiring manager\"\n"
        "    status: VERIFIED\n"
        "    source: test\n"
    )
    app.state.evidence_store = EvidenceStore(profile_dir=str(profile_dir))
    app.state.evidence_checker = EvidenceChecker(
        app.state.evidence_store.verified_values()
    )

    resp = await client.post(f"/api/jobs/{job_id}/prepare")
    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "prepared"
    assert data["tailored_resume"] == "Tailored resume text"
    assert data["cover_letter"] == "Dear hiring manager..."

    mock_tailor.prepare.assert_called_once_with(
        job_description="Build things",
        match_reasons=["Good skills match"],
        suggested_keywords=["Python", "AWS"],
        resume_text=None,
    )

    application = await db.get_application(job_id)
    assert application is not None
    assert application["status"] == "prepared"
    assert application["tailored_resume"] == "Tailored resume text"


@pytest.mark.asyncio
async def test_email_no_job(client):
    resp = await client.post("/api/jobs/999/email")
    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_email_no_cover_letter(client, app):
    db = app.state.db
    job_id = await db.insert_job(
        title="Engineer", company="Acme", location="Remote",
        salary_min=150000, salary_max=200000,
        description="Build things", url="https://example.com/job3",
        posted_date="2026-01-01", application_method="url",
        contact_email="hr@acme.com",
    )
    resp = await client.post(f"/api/jobs/{job_id}/email")
    assert resp.status_code == 400


@pytest.mark.asyncio
async def test_email_success(client, app):
    db = app.state.db
    job_id = await db.insert_job(
        title="Engineer", company="Acme", location="Remote",
        salary_min=150000, salary_max=200000,
        description="Build things", url="https://example.com/job4",
        posted_date="2026-01-01", application_method="email",
        contact_email="hr@acme.com",
    )
    app_id = await db.insert_application(job_id, "prepared")
    await db.update_application(app_id, cover_letter="Dear hiring manager...")

    resp = await client.post(f"/api/jobs/{job_id}/email")
    assert resp.status_code == 200
    data = resp.json()
    assert data["email"]["to"] == "hr@acme.com"
    assert "Engineer" in data["email"]["subject"]


@pytest.mark.asyncio
async def test_email_no_contact(client, app):
    db = app.state.db
    job_id = await db.insert_job(
        title="Engineer", company="Acme", location="Remote",
        salary_min=150000, salary_max=200000,
        description="Build things", url="https://example.com/job5",
        posted_date="2026-01-01", application_method="url",
        contact_email=None,
    )
    app_id = await db.insert_application(job_id, "prepared")
    await db.update_application(app_id, cover_letter="Dear hiring manager...")

    resp = await client.post(f"/api/jobs/{job_id}/email")
    assert resp.status_code == 400


@pytest.mark.asyncio
async def test_get_search_config_empty(client):
    resp = await client.get("/api/search-config")
    assert resp.status_code == 200
    data = resp.json()
    assert data["search_terms"] == []
    assert data["resume_text"] == ""


@pytest.mark.asyncio
async def test_update_search_terms(client, app):
    db = app.state.db
    await db.save_search_config("resume", ["old"])
    resp = await client.post("/api/search-config/terms", json={"search_terms": ["devops remote", "SRE"]})
    assert resp.status_code == 200
    assert resp.json()["search_terms"] == ["devops remote", "SRE"]

    config = await db.get_search_config()
    assert config["search_terms"] == ["devops remote", "SRE"]


@pytest.mark.asyncio
async def test_add_event(client, app):
    db = app.state.db
    job_id = await db.insert_job(
        title="Engineer", company="Acme", location="Remote",
        salary_min=150000, salary_max=200000,
        description="Build things", url="https://example.com/job-event",
        posted_date="2026-01-01", application_method="url",
        contact_email=None,
    )
    resp = await client.post(f"/api/jobs/{job_id}/events", json={"detail": "Looks great"})
    assert resp.status_code == 200
    assert resp.json()["ok"] is True

    resp = await client.get(f"/api/jobs/{job_id}")
    assert resp.status_code == 200
    data = resp.json()
    assert "events" in data
    assert len(data["events"]) == 1
    assert data["events"][0]["event_type"] == "note"
    assert data["events"][0]["detail"] == "Looks great"


@pytest.mark.asyncio
async def test_add_event_with_type(client, app):
    db = app.state.db
    job_id = await db.insert_job(
        title="Engineer", company="Acme", location="Remote",
        salary_min=150000, salary_max=200000,
        description="Build things", url="https://example.com/job-event-type",
        posted_date="2026-01-01", application_method="url",
        contact_email=None,
    )
    call_detail = '{"who": "Jane", "duration": "15 min", "notes": "Discussed role"}'
    resp = await client.post(f"/api/jobs/{job_id}/events", json={"detail": call_detail, "event_type": "call"})
    assert resp.status_code == 200

    resp = await client.get(f"/api/jobs/{job_id}")
    data = resp.json()
    assert data["events"][0]["event_type"] == "call"
    assert "Jane" in data["events"][0]["detail"]


@pytest.mark.asyncio
async def test_add_event_invalid_type(client, app):
    db = app.state.db
    job_id = await db.insert_job(
        title="Engineer", company="Acme", location="Remote",
        salary_min=150000, salary_max=200000,
        description="Build things", url="https://example.com/job-event-bad-type",
        posted_date="2026-01-01", application_method="url",
        contact_email=None,
    )
    resp = await client.post(f"/api/jobs/{job_id}/events", json={"detail": "test", "event_type": "invalid"})
    assert resp.status_code == 400


@pytest.mark.asyncio
async def test_add_event_empty_detail(client, app):
    db = app.state.db
    job_id = await db.insert_job(
        title="Engineer", company="Acme", location="Remote",
        salary_min=150000, salary_max=200000,
        description="Build things", url="https://example.com/job-event-empty",
        posted_date="2026-01-01", application_method="url",
        contact_email=None,
    )
    resp = await client.post(f"/api/jobs/{job_id}/events", json={"detail": ""})
    assert resp.status_code == 400


@pytest.mark.asyncio
async def test_status_change_creates_event(client, app):
    db = app.state.db
    job_id = await db.insert_job(
        title="Engineer", company="Acme", location="Remote",
        salary_min=150000, salary_max=200000,
        description="Build things", url="https://example.com/job-status-event",
        posted_date="2026-01-01", application_method="url",
        contact_email=None,
    )
    resp = await client.post(f"/api/jobs/{job_id}/application?status=applied")
    assert resp.status_code == 200

    events = await db.get_events(job_id)
    assert len(events) == 1
    assert events[0]["event_type"] == "status_change"
    assert "applied" in events[0]["detail"]


@pytest.mark.asyncio
async def test_profile_crud(client):
    # GET should return empty profile
    resp = await client.get("/api/profile")
    assert resp.status_code == 200
    data = resp.json()
    assert data["full_name"] == ""

    # POST to save
    resp = await client.post("/api/profile", json={"full_name": "Test User", "email": "test@x.com"})
    assert resp.status_code == 200

    # GET to verify
    resp = await client.get("/api/profile")
    data = resp.json()
    assert data["full_name"] == "Test User"
    assert data["email"] == "test@x.com"


@pytest.mark.asyncio
async def test_export_csv(client, app):
    from app.database import make_dedup_hash
    db = app.state.db
    dedup = make_dedup_hash("Dev", "Co", "http://x")
    await db.db.execute(
        """INSERT INTO jobs (title, company, location, url, dedup_hash, created_at)
           VALUES (?, ?, ?, ?, ?, ?)""",
        ("Dev", "Co", "Remote", "http://x", dedup, "2026-01-01")
    )
    await db.db.commit()

    resp = await client.get("/api/export/csv")
    assert resp.status_code == 200
    assert "text/csv" in resp.headers["content-type"]
    content = resp.text
    assert "Title" in content
    assert "Dev" in content


@pytest.mark.asyncio
async def test_upload_resume_no_client(client, app):
    app.state._anthropic_client = None
    app.state.testing = True
    import io
    files = {"file": ("resume.txt", io.BytesIO(b"My resume content"), "text/plain")}
    resp = await client.post("/api/resume/upload", files=files)
    assert resp.status_code == 200
    data = resp.json()
    assert data["ok"] is True
    assert data["resume_length"] == len("My resume content")
    assert data["search_terms"] == []


def test_build_ai_client_signature_accepts_settings_object():
    """Regression: second positional arg must be the Settings object, not env_key.

    A router passed `_build_ai_client(ai_settings, env_key)` (old 2-arg form) after
    the signature gained a `settings=` parameter, so a str reached
    `settings.openrouter_api_key` and resume upload 500'd.
    """
    from app.main import _build_ai_client
    from app.config import Settings

    s = Settings(openrouter_api_key="sk-or-test", anthropic_api_key="")
    client = _build_ai_client(None, s, env_key="")
    assert client is not None
    assert client.provider == "openrouter"

    # No sources configured -> no client, and definitely no crash.
    s_empty = Settings(openrouter_api_key="", anthropic_api_key="")
    assert _build_ai_client(None, s_empty, env_key="") is None


@pytest.mark.asyncio
async def test_list_jobs_tolerates_empty_filter_params(client, app):
    """Regression: the Jobs page always sends min_score= (empty) — the endpoint
    must treat empty filters as absent, not 422, or the dashboard shows nothing."""
    resp = await client.get(
        "/api/jobs?limit=25&offset=0&search=&min_score=&sort=score"
        "&work_type=&employment_type=&location=&region=&clearance=&posted_within="
    )
    assert resp.status_code == 200
    assert isinstance(resp.json().get("jobs"), list)


@pytest.mark.asyncio
async def test_upload_resume_appears_in_resumes_list(client, app):
    """Regression: uploading a resume must create a row visible via GET /api/resumes.

    Previously the upload endpoint only wrote search_config, so the onboarding
    checklist kept reporting "No resume yet" after a successful upload."""
    import io
    files = {"file": ("basil_resume.txt", io.BytesIO(b"Basil Ahamed H - AI Engineer resume text"), "text/plain")}
    resp = await client.post("/api/resume/upload", files=files)
    assert resp.status_code == 200
    data = resp.json()
    assert data["ok"] is True
    assert data["resume_id"]

    listing = await client.get("/api/resumes")
    assert listing.status_code == 200
    resumes = listing.json()["resumes"]
    assert len(resumes) == 1
    assert resumes[0]["name"] == "basil_resume.txt"
    assert "Basil Ahamed H" in resumes[0]["resume_text"]
    assert resumes[0]["is_default"] is True  # first upload becomes default

    # Re-uploading the same filename updates in place (no duplicate rows).
    files2 = {"file": ("basil_resume.txt", io.BytesIO(b"Basil Ahamed H - UPDATED resume text"), "text/plain")}
    resp2 = await client.post("/api/resume/upload", files=files2)
    assert resp2.status_code == 200
    listing2 = await client.get("/api/resumes")
    resumes2 = listing2.json()["resumes"]
    assert len(resumes2) == 1
    assert "UPDATED" in resumes2[0]["resume_text"]


@pytest.mark.asyncio
async def test_upload_resume_docx_extracts_text(client, app):
    """Regression: .docx uploads must be parsed as a document, not stored as
    raw ZIP bytes (previously the resume text was 'PK\u0003\u0004...' garbage)."""
    import io
    import docx as docx_lib

    buf = io.BytesIO()
    d = docx_lib.Document()
    d.add_paragraph("Basil Ahamed H")
    d.add_paragraph("AI Engineer with Python and RAG experience")
    d.save(buf)
    buf.seek(0)

    files = {"file": ("basil_cv.docx", buf,
                      "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    resp = await client.post("/api/resume/upload", files=files)
    assert resp.status_code == 200
    data = resp.json()
    assert data["ok"] is True

    listing = await client.get("/api/resumes")
    resumes = listing.json()["resumes"]
    assert len(resumes) == 1
    text = resumes[0]["resume_text"]
    assert text.startswith("PK") is False  # no raw ZIP garbage
    assert "Basil Ahamed H" in text
    assert "RAG experience" in text


@pytest.mark.asyncio
async def test_upload_resume_env_fallback_path_no_crash(client, app, monkeypatch):
    """Regression for the live crash: production path (testing=False, no client on
    state, no DB AI settings) must build-or-skip gracefully, not 500."""
    app.state.testing = False
    app.state.ai_client = None
    if not hasattr(app.state, "settings"):
        from app.config import Settings
        app.state.settings = Settings(openrouter_api_key="", anthropic_api_key="")
    monkeypatch.setattr(app.state.settings, "openrouter_api_key", "", raising=False)
    monkeypatch.setattr(app.state.settings, "anthropic_api_key", "", raising=False)

    import io
    files = {"file": ("resume.txt", io.BytesIO(b"Fallback path resume"), "text/plain")}
    resp = await client.post("/api/resume/upload", files=files)
    assert resp.status_code == 200
    assert resp.json()["ok"] is True


@pytest.mark.asyncio
async def test_upload_resume_pdf(client, app):
    app.state._anthropic_client = None
    app.state.testing = True
    import fitz
    import io
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Senior DevOps Engineer Resume")
    pdf_bytes = doc.tobytes()
    doc.close()
    files = {"file": ("resume.pdf", io.BytesIO(pdf_bytes), "application/pdf")}
    resp = await client.post("/api/resume/upload", files=files)
    assert resp.status_code == 200
    data = resp.json()
    assert data["ok"] is True
    assert data["resume_length"] > 0
